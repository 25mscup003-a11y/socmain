"""Build an editable, 12-slide AJNAT SOC product presentation.

Text, callouts, tables and diagrams are native PowerPoint objects.
Only the actual application screenshots are raster images.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
from lxml import etree

ROOT=Path(__file__).resolve().parents[2]
ASSETS=Path(__file__).parent/'assets'
OUTPUT=ROOT/'AJNAT_SOC_Brochure_Improved.pptx'
P=Presentation()
P.slide_width=Inches(13.333333)
P.slide_height=Inches(7.5)
P.core_properties.title='AJNAT SOC | Product Overview & UI Walkthrough'
P.core_properties.subject='Visibility, investigation and coordinated response'
P.core_properties.author='AJNAT SOC'
P.core_properties.keywords='AJNAT, SOC, SIEM, EDR, SOAR, product overview'
P.core_properties.comments='Actual repository UI screenshots use clearly labelled demonstration data.'
C={'navy':'091D2A','dark':'071622','panel':'102B3A','teal':'008C79','mint':'64DDC1',
   'lime':'D6F38D','white':'FFFFFF','paper':'F4F6F3','ink':'132F3D','muted':'526875',
   'line':'D4DFDB','soft':'E6EFEB','pale':'B5C9D0','gold':'F2BE68'}
NOTES=[]

def color(c): return RGBColor.from_string(C.get(c,c))
def rect(s,x,y,w,h,fill='white',line=None,r=0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
    sh.line.fill.background() if line is None else None
    if line: sh.line.color.rgb=color(line);sh.line.width=Pt(.8)
    if r: sh.adjustments[0]=r
    return sh

def txt(s,t,x,y,w,h,size=18,c='ink',bold=False,font='Lato',align=PP_ALIGN.LEFT):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_top=0;tf.margin_bottom=0;tf.margin_left=0;tf.margin_right=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.08
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c)
    return tb

def line(s,x1,y1,x2,y2,c='line',width=1):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=color(c);sh.line.width=Pt(width);return sh

def dot(s,n,x,y,fill='teal',fg='white',d=.34):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.fill.background()
    tf=sh.text_frame;tf.clear();tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0];p.text=str(n);p.alignment=PP_ALIGN.CENTER
    p.font.name='Lato';p.font.size=Pt(12);p.font.bold=True;p.font.color.rgb=color(fg)
    return sh

def pill(s,t,x,y,w,fill='soft',fg='teal'):
    rect(s,x,y,w,.29,fill,r=.16)
    txt(s,t,x,y+.054,w,.18,9,fg,True,align=PP_ALIGN.CENTER)

def base(section,dark=False):
    s=P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=color('navy' if dark else 'paper')
    txt(s,'AJNAT',.52,.29,1.1,.26,17,'white' if dark else 'ink',True)
    line(s,1.75,.31,1.75,.55,'panel' if dark else 'line')
    txt(s,'SOC  /  '+section.upper(),1.93,.35,8,.18,9,'pale' if dark else 'muted',True)
    line(s,.52,6.99,12.81,6.99,'panel' if dark else 'line')
    txt(s,'AJNAT SOC  ·  VISIBILITY. CONTEXT. CONTROL.',.52,7.15,8,.15,8,'pale' if dark else 'muted')
    txt(s,f'{len(P.slides):02d} / 12',11.98,7.12,.83,.19,9,'pale' if dark else 'muted',align=PP_ALIGN.RIGHT)
    return s

def heading(s,kicker,title,sub=None,dark=False):
    txt(s,kicker.upper(),.55,.87,11,.2,10,'mint' if dark else 'teal',True)
    txt(s,title,.52,1.23,12.1,.66,31,'white' if dark else 'ink',True)
    if sub: txt(s,sub,.55,1.98,11.8,.53,16,'pale' if dark else 'muted')

def note(s,title,hinglish,source=None):
    extra=''
    if source:
      extra='\n\nVisual provenance: Actual AJNAT React UI from '+source+'. Captured locally with browser-intercepted demonstration API responses; no production telemetry. UI labels and sample counts do not establish measured product performance. See docs/brochure/capture_ui.py and assets/capture-manifest.json.'
    s.notes_slide.notes_text_frame.text=title+'\n\nPresenter notes (Hinglish):\n'+hinglish+extra
    NOTES.append((title,hinglish,source))

def image_contain(s,path,x,y,w,h):
    im=Image.open(path);iw,ih=im.size;factor=min(w/iw,h/ih);nw,nh=iw*factor,ih*factor
    s.shapes.add_picture(str(path),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))
    return x+(w-nw)/2,y+(h-nh)/2,nw,nh

def screenshot(s,name,x=.55,y=2.08,w=8.85,h=4.42):
    rect(s,x-.045,y-.045,w+.09,h+.09,'navy',r=.025)
    # Screenshots are cropped, never redrawn or altered to imitate product UI.
    return image_contain(s,ASSETS/f'{name}-slide.png',x,y,w,h)

def callout(s,n,title,body,y,x=9.76,w=2.99):
    dot(s,n,x,y)
    txt(s,title,x,y+.46,w,.33,17,'ink',True)
    txt(s,body,x,y+.89,w,.71,13.5,'muted')

def ui_slide(kicker,title,sub,img,items,note_text,source):
    s=base('Product walkthrough')
    txt(s,kicker.upper(),.55,.87,11,.2,10,'teal',True)
    txt(s,title,.52,1.21,12.1,.53,29,'ink',True)
    txt(s,sub,.55,1.81,12,.28,13.5,'muted')
    screenshot(s,img,y=2.28,h=4.26)
    for i,(t,b) in enumerate(items): callout(s,i+1,t,b,2.34+i*1.45)
    txt(s,'ACTUAL AJNAT UI  /  SAMPLE DATA',.56,6.69,8.8,.16,8.5,'muted',True)
    note(s,title,note_text,source)
    return s

def prepare_images():
    # Preserve product geometry; crops focus on the parts discussed on each slide.
    for name in ['dashboard','correlation','investigation','soar','playbooks','edr']:
      src=ASSETS/f'{name}.png'
      im=Image.open(src)
      width,height=im.size
      if name=='edr':
        # Focus on monitoring cards; omit the live-telemetry header in this demo capture.
        box=(25,330,width-25,1035)
      elif name=='dashboard':
        box=(0,138,width,1018)
      elif name=='investigation':
        box=(0,304,width,min(height,1236))
      elif name=='soar':
        box=(25,335,width-25,1135)
      elif name=='correlation':
        box=(0,15,width,955)
      else:
        box=(0,0,width,min(height,int(width/2.078)))
      im.crop(box).save(ASSETS/f'{name}-slide.png')

prepare_images()

# 01 — Brand and promise
s=base('Product overview',True)
pill(s,'UNIFIED SECURITY OPERATIONS',.56,1.16,3.12,'panel','mint')
txt(s,'AJNAT SOC',.52,1.86,6.2,.79,48,'white',True)
txt(s,'See the signals.\nUnderstand the threat.\nCoordinate the response.',.56,2.88,6.3,1.83,30,'white',True)
txt(s,'Endpoint visibility, connected evidence and\nincident workflows in one analyst workspace.',.58,5.01,5.8,.76,17,'pale')
for i,(a,w) in enumerate([('SIEM',1.12),('EDR',1.1),('SOAR',1.22)]): pill(s,a,.58+i*1.4,6.2,w,'panel','mint')
rect(s,7.17,1.17,5.59,5.51,'panel',r=.04)
txt(s,'FROM SIGNAL TO ACTION',7.5,1.52,4.9,.25,12,'mint',True)
image_contain(s,ASSETS/'correlation-slide.png',7.42,2.08,5.12,2.66)
line(s,7.51,5.13,12.39,5.13,'pale',.8)
txt(s,'Visibility',7.51,5.52,1.5,.35,18,'white',True)
txt(s,'Context',9.24,5.52,1.5,.35,18,'white',True)
txt(s,'Control',10.97,5.52,1.5,.35,18,'white',True)
txt(s,'ACTUAL AJNAT UI · SAMPLE DATA',7.51,6.21,4.8,.18,8.5,'pale')
note(s,'AJNAT SOC — Product overview','AJNAT SOC security team ko ek common workspace deta hai. Endpoint, network aur identity ke signals ko saath dekhkar analyst threat samajh sakta hai aur response coordinate kar sakta hai. Aage hum actual UI screens aur ek short incident example dekhenge. Screenshots mein sample data hai.', 'company/src/pages/CorrelationPage.jsx')

# 02 — Simple product definition
s=base('What it does')
heading(s,'The platform','One workspace for the security team.','A SOC (Security Operations Center) monitors activity, investigates threats and coordinates response.')
cards=[('SIEM','Bring events together','Security Information and\nEvent Management','Collect and connect security logs\nfrom configured sources.'),
       ('EDR','Understand the endpoint','Endpoint Detection\nand Response','Review process, file, network\nand authentication activity.'),
       ('SOAR','Coordinate the next action','Security Orchestration,\nAutomation and Response','Use rules, playbooks and approvals\nto manage supported actions.')]
for i,(abbr,title,full,body) in enumerate(cards):
    x=.55+i*4.16
    rect(s,x,2.79,3.93,3.22,'white',r=.035)
    rect(s,x,2.79,3.93,.065,'teal')
    txt(s,abbr,x+.23,3.03,3.4,.47,29,'teal',True)
    txt(s,full,x+.23,3.62,3.45,.53,12,'muted')
    txt(s,title,x+.23,4.41,3.45,.35,19,'ink',True)
    txt(s,body,x+.23,4.99,3.4,.8,15,'muted')
txt(s,'Connected sources',.58,6.43,1.9,.27,12,'teal',True)
txt(s,'Endpoints  /  Network  /  Identity  /  Web applications',2.57,6.4,9.9,.3,16,'ink')
note(s,'One workspace for the security team','SOC ka simple meaning hai security monitoring aur incident response ka central function. SIEM logs ko collect aur correlate karta hai. EDR endpoint ke andar ki activity dikhata hai. SOAR approved workflows se response coordinate karta hai. AJNAT in capabilities ko ek operational context mein laata hai.')

# 03 — Editable process diagram
s=base('How it works',True)
heading(s,'The operating flow','From raw events to a clear next step.','Each stage adds context so an analyst can make an informed decision.',True)
steps=[('01','Collect','Endpoint, network,\nweb and login events','Shared event visibility'),
       ('02','Correlate','Related activity,\nentities and indicators','A connected attack story'),
       ('03','Investigate','Severity, timeline,\nevidence and ownership','An analyst assessment'),
       ('04','Respond','Approvals, supported\nactions and execution logs','A recorded outcome')]
for i,(num,title,body,out) in enumerate(steps):
    x=.55+i*3.13
    rect(s,x,2.93,2.86,2.93,'panel',r=.025)
    txt(s,num,x+.23,3.18,2.3,.55,32,'mint',True)
    txt(s,title,x+.23,3.99,2.3,.4,23,'white',True)
    txt(s,body,x+.23,4.6,2.4,.76,16,'pale')
    txt(s,out,x+.02,6.1,2.8,.44,12,'mint',True,align=PP_ALIGN.CENTER)
    if i<3: txt(s,'→',x+2.88,4.11,.23,.4,19,'mint',True)
note(s,'From raw events to a clear next step','Pehle configured sources se events aate hain. Correlation related activity ko link karta hai. Analyst severity, timeline aur evidence review karke incident assess karta hai. Uske baad authorized response actions, approvals aur execution result case ke saath record hote hain. Har alert automatically confirmed attack nahi hota.')

# 04–08 — Real product UI with short explanations
ui_slide('01 / SOC overview','Start with operational priorities.','A manager view brings company coverage, work queues and security activity together.','dashboard',
    [('Coverage','Review assigned companies and security modules.'),
     ('Work queues','See alerts, incidents and escalations in one view.'),
     ('Follow-up','Open reports and audit views to support review and handover.')],
    'Yeh actual SOC Manager dashboard hai. Ismein assigned companies, IDS/IPS, firewall aur operational queues ka overview milta hai. Manager yahan se relevant module ya queue open kar sakta hai. Screen ke numbers sirf demonstration data hain; performance ya customer results nahi.',
    'company/src/pages/analyst/soc/RoleDashboardPage.jsx')

ui_slide('02 / Endpoint visibility','Understand activity on the endpoint.','Monitoring modules organize device activity into focused investigation views.','edr',
    [('Processes & files','Inspect execution activity and file changes.'),
     ('Network & identity','Review connections, DNS and authentication signals.'),
     ('Focused coverage','Choose modules for the endpoint platform and enabled sensors.')],
    'EDR page endpoint activity ko monitoring modules mein organize karta hai: process, file, network, authentication aur doosre security areas. Analyst required module open karke related telemetry inspect kar sakta hai. Coverage operating system, permissions, sensors aur configuration par depend karti hai. Shown counters are sample data; this is a cropped view of the existing monitoring-card grid.',
    'company/src/pages/EDRPage.jsx')

ui_slide('03 / Event correlation','Connect signals into an attack story.','Related events appear together, with the affected host, severity and supporting evidence.','correlation',
    [('Pattern','See which sequence of activity triggered the correlation.'),
     ('Context','Review the endpoint, risk and investigation status.'),
     ('Evidence','Open linked events before confirming the incident.')],
    'Correlation page alag events ko ek possible attack chain ke andar dikhata hai. Sample case mein failed logins, PowerShell activity aur outbound connection FIN-WS-014 se linked hain. Analyst pattern, risk aur linked evidence check karta hai. Confidence score ek signal hai; final conclusion evidence review ke baad hota hai.',
    'company/src/pages/CorrelationPage.jsx')

ui_slide('04 / Investigation','Read the evidence before taking action.','Follow the event timeline and inspect the details behind a suspicious signal.','investigation',
    [('Find','Filter by severity, source, status or an indicator.'),
     ('Inspect','Review host, user, process and network details.'),
     ('Assess','Use the timeline, raw logs and analyst notes for the decision.')],
    'Investigation screen mein left side par events hain aur selected event ki detail right side par aati hai. Analyst host, user, process aur source/destination details ko inspect kar sakta hai. Timeline, correlated events, IOC intelligence, raw logs aur notes context dete hain. AI summaries jahan configured hon wahan assistance de sakti hain; analyst review zaroori hai.',
    'company/src/pages/ThreatInvestigationPage.jsx')

ui_slide('05 / Coordinated response','Move from review to a managed response.','SOAR brings automation rules, execution history and pending approvals into one workspace.','soar',
    [('Orchestrate','Use rules and playbooks to organize response steps.'),
     ('Approve','Review sensitive actions before they execute.'),
     ('Track','Keep the execution status and outcome with the workflow.')],
    'SOAR ka role response steps ko organize karna hai. Rules aur playbooks incident creation, assignment aur supported containment actions coordinate kar sakte hain. Pending approval queue sensitive actions par human control deti hai. Endpoint isolation ya blocking permissions, platform aur integrations par depend karta hai. Screen ke timings aur counts sample values hain, benchmarks nahi.',
    'company/src/pages/SoarPage.jsx')

# 09 — Short, concrete scenario
s=base('Incident example',True)
heading(s,'Illustrative scenario','One endpoint. Three connected signals.','FIN-WS-014  /  Possible account compromise  /  Requires analyst confirmation',True)
events=[('01','Failed logins','Repeated authentication\nfailures for finance.user.'),
        ('02','Suspicious execution','PowerShell activity on\nthe same endpoint.'),
        ('03','Outbound connection','A destination that needs\nindicator enrichment.')]
for i,(n,t,b) in enumerate(events):
    x=.55+i*4.19
    rect(s,x,2.81,3.93,2.05,'panel',r=.03)
    dot(s,n,x+.23,3.04,'mint','navy',.4)
    txt(s,t,x+.23,3.64,3.45,.4,22,'white',True)
    txt(s,b,x+.23,4.17,3.4,.6,15,'pale')
    if i<2: txt(s,'→',x+3.96,3.64,.23,.35,18,'mint',True)
line(s,.79,5.2,12.55,5.2,'mint',1.2)
for i,(t,b) in enumerate([('Connect the evidence','Correlate the host, identity and time window.'),
                        ('Review the case','Confirm scope, severity and incident ownership.'),
                        ('Choose the response','Approve supported action and record its result.')]):
    x=.57+i*4.19
    txt(s,t,x,5.62,3.9,.35,19,'mint',True)
    txt(s,b,x,6.12,3.8,.54,14,'pale')
note(s,'One endpoint. Three connected signals','Is example mein ek endpoint par teen signals milte hain: repeated failed logins, suspicious PowerShell execution aur unusual outbound connection. AJNAT in events ko link karne aur evidence inspect karne ka workflow deta hai. Analyst scope aur severity confirm karta hai, owner assign hota hai, phir supported response ko approve aur track kiya ja sakta hai. Yeh illustrative case hai; actual breach ya measured outcome nahi.')

# 10 — Role ownership and auditable work
s=base('Operational control')
heading(s,'Teams & accountability','Clear ownership at every stage.','Role-based workspaces keep investigation, escalation and follow-up organized.')
roles=[('SOC manager','Priorities, assignment\nand team oversight.'),
       ('L1 analyst','Initial review, ticket triage\nand escalation.'),
       ('L2 / L3 analysts','Deeper investigation\nand response coordination.'),
       ('L4 threat intelligence','Indicator context and\nthreat intelligence work.')]
for i,(t,b) in enumerate(roles):
    x=.55+i*3.13
    rect(s,x,2.78,2.86,1.77,'white',r=.04)
    txt(s,t,x+.2,3.03,2.45,.5,19,'ink',True)
    txt(s,b,x+.2,3.68,2.5,.65,14,'muted')
rect(s,.55,4.91,12.24,1.66,'navy',r=.03)
for i,(t,b) in enumerate([('Scoped access','Company and role context'),('Case history','Owner, decisions and actions'),('Reports & audit','Evidence for review and handover')]):
    x=.8+i*4.06
    txt(s,t,x,5.22,3.52,.35,20,'mint',True)
    txt(s,b,x,5.83,3.64,.46,14,'white')
note(s,'Clear ownership at every stage','SOC manager team priorities aur assignments manage karta hai. L1 initial triage karta hai; L2/L3 detailed investigation aur response coordination karte hain. L4 threat intelligence context par kaam karta hai. Company scope aur role permissions determine karte hain ki user kya dekh aur kar sakta hai. Case history, audit aur reports handover aur review ko support karte hain. Yeh compliance certification ka claim nahi hai.')

# 11 — Benefits without unsupported numbers
s=base('Why it matters')
heading(s,'Operational value','Less switching. More connected context.','Keep the investigation and the response in the same operational conversation.')
rect(s,.55,2.73,12.24,.53,'navy')
txt(s,'FRAGMENTED OPERATIONS',.8,2.88,5.42,.2,11,'pale',True)
txt(s,'WITH AJNAT SOC',6.87,2.88,5.45,.2,11,'mint',True)
rows=[('Separate asset and event views','Shared visibility across connected sources'),
      ('Manual reconstruction of related activity','Correlated events with investigation context'),
      ('Indicators checked in separate tools','Threat intelligence alongside relevant evidence'),
      ('Response decisions spread across handoffs','Owned cases, approvals and execution records')]
for i,(a,b) in enumerate(rows):
    y=3.26+i*.65
    rect(s,.55,y,12.24,.65,'white' if i%2==0 else 'soft')
    txt(s,a,.8,y+.18,5.25,.36,15,'muted')
    txt(s,'→',6.33,y+.16,.3,.3,16,'teal',True)
    txt(s,b,6.87,y+.18,5.51,.36,15,'ink',True)
for i,(a,b) in enumerate([('VISIBILITY','Know what is happening'),('CONTEXT','Understand what matters'),('CONTROL','Choose the next action')]):
    x=.58+i*4.15
    txt(s,a,x,6.18,3.89,.19,10,'teal',True)
    txt(s,b,x,6.52,3.89,.3,17,'ink',True)
note(s,'Less switching. More connected context','Connected workflow ka benefit hai ki analyst ko assets, events, indicators aur response context ek jagah milta hai. Isse review aur handover structured ho sakte hain. Is slide mein operational benefits explain kiye gaye hain; time saving, detection rate ya ROI ke unverified percentage claims nahi diye gaye.')

# 12 — Concrete next steps
s=base('Pilot & next steps',True)
txt(s,'PUT IT IN CONTEXT',.56,1.04,11,.23,11,'mint',True)
txt(s,'See AJNAT SOC\nin your environment.',.53,1.69,11.8,1.65,43,'white',True)
txt(s,'Start with your critical assets, relevant telemetry and a realistic incident workflow.',.59,3.55,11.2,.63,19,'pale')
for i,(n,t,b) in enumerate([('01','Define the scope','Choose endpoints, sources\nand access roles.'),
                          ('02','Walk through a case','Follow detection, investigation\nand an approved response.'),
                          ('03','Agree on success','Review coverage, evidence quality\nand handover readiness.')]):
    x=.59+i*4.17
    line(s,x,4.6,x+3.87,4.6,'panel',1.8)
    txt(s,n,x,4.92,.62,.42,25,'mint',True)
    txt(s,t,x+.77,4.96,3.03,.41,22,'white',True)
    txt(s,b,x+.77,5.6,3.05,.81,16,'pale')
note(s,'See AJNAT SOC in your environment','Next step ek focused pilot hai. Pehle critical assets, data sources aur user roles select karein. Phir ek realistic incident ko detection se investigation aur approved response tak walk through karein. Success criteria coverage, evidence quality, action support aur team handover ke around define karein. Deployment se pehle platform support, integrations aur agent permissions confirm karna hoga.')

assert len(P.slides)==12
# Keep the presentation flat and clean instead of inheriting theme shadows.
for slide in P.slides:
    for shape in slide.shapes:
        sppr=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if sppr is not None:
            etree.SubElement(sppr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
P.save(OUTPUT)
md=['# AJNAT SOC — Presenter Notes','', 'Slides use concise English. Notes below provide a short Hinglish speaking guide.', '']
for i,(title,n,src) in enumerate(NOTES,1):
    md += [f'## {i:02d}. {title}','',n,'']
    if src: md += [f'Visual: `{src}`. Actual UI, sample data.','']
(Path(__file__).parent/'Presenter_Notes.md').write_text('\n'.join(md))
print(f'Created {OUTPUT} ({len(P.slides)} slides)')
