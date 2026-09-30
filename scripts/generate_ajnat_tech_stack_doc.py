from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "AJNAT_SOC_Technology_Stack.docx"
NAVY = "102A43"
BLUE = "1677FF"
CYAN = "00A6A6"
LIGHT_BLUE = "EAF3FF"
LIGHT_GREY = "F3F6F9"
MID_GREY = "627D98"
WHITE = "FFFFFF"
GREEN = "168A5B"
AMBER = "B7791F"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=NAVY, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True, WHITE, 9)
        shade(table.rows[0].cells[i], NAVY)
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, False, NAVY, 8.5)
            if row_idx % 2 == 0:
                shade(cells[i], LIGHT_GREY)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_callout(doc, title, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(0.12)
    table.columns[1].width = Inches(6.85)
    shade(table.cell(0, 0), accent)
    shade(table.cell(0, 1), fill)
    table.cell(0, 0).text = ""
    cell = table.cell(0, 1)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.color.rgb = RGBColor.from_string(NAVY)
    r2.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, color in [("Title", 30, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 10.5, CYAN)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].paragraph_format.space_before = Pt(8)
styles["Heading 2"].paragraph_format.space_after = Pt(5)

for sec in doc.sections:
    header = sec.header.paragraphs[0]
    header.text = "AJNAT SOC  /  TECHNOLOGY STACK"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)
    add_page_number(sec.footer.paragraphs[0])

# Cover page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(70)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AJNAT")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(24)
r.font.color.rgb = RGBColor.from_string(BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
r = p.add_run("SOC PLATFORM")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(34)
r.font.color.rgb = RGBColor.from_string(NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
r = p.add_run("Technology Stack, Cloud Architecture & Security Overview")
r.font.name = "Aptos Display"
r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(CYAN)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for idx, (label, value) in enumerate([
    ("DOCUMENT", "Technical Overview"),
    ("STATUS", "Current + Target State"),
    ("DATE", date(2026, 9, 7).strftime("%d %B %Y")),
]):
    cell = table.cell(0, idx)
    shade(cell, LIGHT_BLUE)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(label + "\n").bold = True
    run = p.add_run(value)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)

doc.add_paragraph("\n")
add_callout(
    doc,
    "Purpose",
    "A management-friendly and engineering-ready view of the technologies used by AJNAT SOC, the current deployment posture, and the recommended cloud production architecture.",
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(85)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prepared from the AJNAT SOC source repository and architecture documentation")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(MID_GREY)

doc.add_page_break()

add_heading(doc, "1. Executive Summary", 1)
doc.add_paragraph(
    "AJNAT SOC is a multi-tenant security operations platform that combines web dashboards, a Node.js control plane, endpoint agents, network-security sensors, event streaming, and analytics storage. Its architecture supports real-time monitoring, alerting, endpoint visibility, IDS/IPS, threat intelligence, and automated response workflows."
)

add_table(doc, ["Layer", "Primary Technology", "Role"], [
    ("Web applications", "React 18 + Vite 5", "Company and super-admin dashboards"),
    ("API & real-time", "Node.js + Express 4 + Socket.IO", "REST APIs, authentication, live telemetry and commands"),
    ("Operational database", "MongoDB + Mongoose 7", "Tenants, users, agents, configuration and alerts"),
    ("Event pipeline", "Kafka / KafkaJS", "Scalable ingestion and asynchronous processing"),
    ("Analytics", "ClickHouse support", "High-volume event search and long-term analytics"),
    ("Cache & coordination", "Redis", "Caching, sessions, queues and distributed coordination"),
    ("Endpoint security", "Python agent + Android Java agent", "Host telemetry, monitoring and response"),
    ("Security sensors", "Suricata, Zeek, YARA, VirusTotal, Velociraptor", "Network and endpoint detection/enrichment"),
], [1.25, 2.05, 3.65])

add_callout(
    doc,
    "Important deployment note",
    "The repository contains a local Kafka Docker Compose stack and detailed target-cloud architecture. It does not yet contain a complete, reproducible production Docker/Kubernetes/CI deployment for the full platform. This distinction is important when presenting current readiness.",
    fill="FFF8E6",
    accent=AMBER,
)

add_heading(doc, "2. Platform Architecture", 1)
add_heading(doc, "2.1 Logical data flow", 2)
flow = doc.add_table(rows=1, cols=7)
flow.alignment = WD_TABLE_ALIGNMENT.CENTER
flow.autofit = False
labels = ["Endpoints &\nSensors", "→", "API /\nIngestion", "→", "Kafka &\nWorkers", "→", "SOC UI /\nAnalytics"]
for i, label in enumerate(labels):
    cell = flow.cell(0, i)
    cell.width = Inches(0.95 if i % 2 == 0 else 0.25)
    if i % 2 == 0:
        shade(cell, NAVY if i in (0, 6) else BLUE)
        set_cell_text(cell, label, True, WHITE, 9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        set_cell_text(cell, label, True, CYAN, 14)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

add_bullets(doc, [
    "Python/Linux and Android agents collect endpoint, process, file, USB, memory, network, DNS and VPN activity.",
    "Suricata and Zeek provide network detection and protocol visibility; YARA and VirusTotal enrich malware analysis.",
    "The backend authenticates agents and users, validates payloads, persists operational data and emits live updates through Socket.IO.",
    "Kafka decouples ingestion from downstream processing; Redis and ClickHouse support high-throughput scale patterns.",
    "React dashboards provide tenant operations, security monitoring, administration and investigation views.",
])

add_heading(doc, "2.2 Major applications", 2)
add_table(doc, ["Application", "Audience", "Primary Responsibility"], [
    ("Company portal", "SOC analysts and tenant administrators", "Monitoring, alerts, assets, agents, investigations and policy"),
    ("Super-admin portal", "Platform administrators", "Tenant lifecycle, platform-wide administration and oversight"),
    ("Backend services", "Applications, agents and integrations", "API, identity, telemetry, orchestration and real-time messaging"),
    ("SOC Agent", "Linux/endpoint systems", "Continuous security telemetry, local monitoring and response"),
    ("Android Agent", "Managed Android devices", "Mobile endpoint visibility and policy enforcement"),
], [1.35, 2.0, 3.6])

doc.add_page_break()
add_heading(doc, "3. Detailed Technology Stack", 1)

add_heading(doc, "3.1 Frontend", 2)
add_table(doc, ["Technology", "Version / Family", "Use in AJNAT SOC"], [
    ("React", "18.x", "Component-based company and super-admin interfaces"),
    ("Vite", "5.x", "Development server and production bundling"),
    ("React Router", "6.x", "Client-side routing and protected application flows"),
    ("Axios", "1.6+", "REST API communication"),
    ("Socket.IO Client", "4.6+", "Live alerts, agent status and real-time updates"),
    ("Leaflet", "1.9+", "Geospatial and map-based visualization"),
], [1.45, 1.45, 4.05])

add_heading(doc, "3.2 Backend and integration", 2)
add_table(doc, ["Technology", "Version / Family", "Use in AJNAT SOC"], [
    ("Node.js", "JavaScript runtime", "Backend execution environment"),
    ("Express", "4.18+", "REST APIs, middleware and routing"),
    ("Mongoose", "7.x", "MongoDB schema and data access layer"),
    ("Socket.IO", "4.6+", "Bidirectional real-time communication"),
    ("JWT", "jsonwebtoken 9.x", "User and service authentication tokens"),
    ("KafkaJS", "2.2+", "Kafka producer/consumer integration"),
    ("AWS SDK", "S3 + KMS clients", "Object storage and cloud key-management integration"),
], [1.45, 1.45, 4.05])

add_heading(doc, "3.3 Data and messaging", 2)
add_table(doc, ["Technology", "Workload", "Design Position"], [
    ("MongoDB", "Operational data and configuration", "Current primary database"),
    ("Redis", "Cache, ephemeral state and coordination", "Integrated backend capability"),
    ("Kafka", "Durable event streaming", "Local Compose stack and target production pipeline"),
    ("ClickHouse", "Security-event analytics", "Supported/target high-volume analytics store"),
    ("S3-compatible storage", "Artifacts, exports and cold data", "Target cloud object-storage layer"),
], [1.5, 2.35, 3.1])

add_heading(doc, "3.4 Endpoint agents", 2)
add_table(doc, ["Agent", "Technology", "Core Capabilities"], [
    ("SOC endpoint agent", "Python 3", "Host inventory, processes, connections, FIM, network-change detection, security telemetry and response"),
    ("Python libraries", "psutil, watchdog, cryptography, yara-python, pyudev", "System telemetry, file events, encryption, malware rules and Linux device monitoring"),
    ("Android agent", "Java 17; Android SDK 35; minimum SDK 23", "Managed mobile visibility, telemetry and device-side controls"),
], [1.45, 2.1, 3.4])

doc.add_page_break()
add_heading(doc, "4. Security Technology Stack", 1)
add_table(doc, ["Control Area", "Technology / Mechanism", "Purpose"], [
    ("Payload encryption", "AES-256-GCM", "Confidentiality and authenticated encryption for sensitive agent data"),
    ("Transport", "TLS; mTLS-capable architecture", "Encrypted service communication and strong machine identity"),
    ("Authentication", "JWT and agent credentials", "Authenticated users, sessions and endpoint agents"),
    ("Password security", "Argon2 / bcrypt", "Strong one-way credential hashing"),
    ("Agent trust", "Signed payloads and integrity manifests", "Tamper detection and controlled agent updates"),
    ("Endpoint detection", "YARA, process/memory/file monitoring", "Malware and suspicious-behavior detection"),
    ("Network detection", "Suricata IDS/IPS + Zeek", "Signature detection, traffic metadata and protocol analysis"),
    ("Threat intelligence", "VirusTotal and enrichment workflows", "Reputation and indicator context"),
    ("Investigation", "Velociraptor integration", "Endpoint query and incident-response support"),
], [1.4, 2.25, 3.3])

add_heading(doc, "4.1 Security design principles", 2)
add_bullets(doc, [
    "Tenant isolation must be enforced in every database query, event stream and object-storage path.",
    "Secrets and encryption keys should be stored in a managed KMS/secret manager, never in images or source code.",
    "Agent updates should retain integrity verification, signed manifests and rollback capability.",
    "Audit logs should be immutable, time-synchronized and retained independently from application logs.",
    "Network egress should be allow-listed for agents, sensors, enrichments and platform services.",
])

add_heading(doc, "5. Cloud and Production Architecture", 1)
add_heading(doc, "5.1 Recommended deployment model", 2)
add_table(doc, ["Cloud Layer", "Recommended Service Pattern", "AJNAT Workload"], [
    ("Edge", "DNS, CDN/WAF, DDoS protection, load balancer", "Public web and API entry points"),
    ("Application", "Stateless containers on Kubernetes or managed container service", "API, Socket.IO gateways and workers"),
    ("Streaming", "Managed Kafka-compatible broker", "Telemetry ingestion and processing topics"),
    ("Operational data", "Managed MongoDB with backups and replicas", "Tenant, identity, policy and operational records"),
    ("Analytics", "Managed/clustered ClickHouse", "High-volume security events and fast investigations"),
    ("Cache", "Managed Redis with replication", "Caching and distributed coordination"),
    ("Objects", "S3-compatible versioned storage", "Evidence, files, exports and archived data"),
    ("Keys & secrets", "Cloud KMS + managed secret store", "Encryption keys, certificates and credentials"),
    ("Observability", "Central logs, metrics, traces and alerting", "Service health, latency, errors and capacity"),
], [1.25, 2.55, 3.15])

add_heading(doc, "5.2 Cloud controls", 2)
add_bullets(doc, [
    "Private subnets for databases, brokers, workers and administrative services.",
    "Horizontal autoscaling based on API latency, worker lag, Kafka consumer lag and queue depth.",
    "Multi-zone deployment for critical APIs, brokers, databases and Redis.",
    "Encrypted backups with tested restore procedures and defined recovery objectives.",
    "Infrastructure as Code, signed container images, dependency scanning and gated releases.",
])

doc.add_page_break()
add_heading(doc, "6. Current State vs Target State", 1)
add_table(doc, ["Capability", "Current Repository Evidence", "Target Production State"], [
    ("Web/API", "React applications and Express backend implemented", "Containerized stateless services behind managed load balancing"),
    ("Database", "MongoDB/Mongoose operational model", "Managed replicated MongoDB with backups and tenant-aware indexes"),
    ("Streaming", "KafkaJS integration and local Kafka Compose", "Managed multi-zone Kafka with retention, ACLs and monitoring"),
    ("Analytics", "ClickHouse support documented", "Production cluster with lifecycle tiers and query governance"),
    ("Cache", "Redis integration present", "Highly available managed Redis"),
    ("Agents", "Python and Android implementations", "Signed, staged and observable release channels"),
    ("Cloud", "AWS S3/KMS libraries and target architecture", "Infrastructure-as-Code production environment"),
    ("Delivery", "No complete full-platform production CI/CD found", "Automated build, test, scan, deploy and rollback pipeline"),
], [1.25, 2.8, 2.9])

add_callout(
    doc,
    "Interpretation",
    "AJNAT SOC has a broad implemented application and security stack. The largest remaining production-engineering gap is repeatable full-platform deployment: container standards, Kubernetes/managed-service definitions, CI/CD, infrastructure code, observability baselines and disaster-recovery automation.",
    fill="EAF9F3",
    accent=GREEN,
)

add_heading(doc, "7. Recommended Delivery Roadmap", 1)
add_numbered(doc, [
    "Baseline the platform: publish supported Node.js, Python, Java, MongoDB, Redis, Kafka and ClickHouse versions; lock dependencies and define environment contracts.",
    "Containerize services: create hardened, non-root images for backend, web applications and workers; generate an SBOM for every release.",
    "Codify cloud infrastructure: provision network, Kubernetes/containers, broker, databases, KMS, secrets, object storage and observability through Infrastructure as Code.",
    "Build release gates: unit/integration tests, SAST, dependency and image scanning, signed artifacts, agent-manifest verification and automated rollback.",
    "Validate scale and resilience: load-test ingestion, consumer lag, Socket.IO fan-out, MongoDB indexes and ClickHouse queries; run backup-restore and zone-failure exercises.",
    "Operationalize security: formalize key rotation, certificate lifecycle, tenant isolation tests, immutable audit retention and incident-response runbooks.",
])

add_heading(doc, "8. Suggested Ownership", 1)
add_table(doc, ["Domain", "Primary Owner", "Key Deliverables"], [
    ("Application platform", "Backend and frontend engineering", "APIs, portals, real-time services and application tests"),
    ("Detection engineering", "SOC / security engineering", "Suricata, Zeek, YARA, enrichment and response content"),
    ("Endpoint agents", "Agent engineering", "Secure collection, durable delivery, update integrity and compatibility"),
    ("Cloud platform", "DevOps / SRE", "Infrastructure, CI/CD, observability, scaling and disaster recovery"),
    ("Security assurance", "Product security", "Threat models, release gates, secrets, identity and audit controls"),
], [1.45, 2.05, 3.45])

doc.add_page_break()
add_heading(doc, "9. Repository Evidence", 1)
doc.add_paragraph("This document was prepared from the following project files:")
add_table(doc, ["Source", "Evidence Used"], [
    ("backend/package.json", "Backend framework, database, messaging, Redis, AWS and security dependencies"),
    ("company/package.json", "Company portal frontend stack"),
    ("superadmin/package.json", "Super-admin frontend stack"),
    ("backend/soc-agent/requirements.txt", "Python endpoint-agent dependencies and capabilities"),
    ("backend/android-agent/app/build.gradle", "Android Java, SDK and compatibility levels"),
    ("kafka/docker-compose.yml", "Local Kafka development deployment"),
    ("docs/architecture/current-architecture.md", "Current topology and application/data flow"),
    ("docs/architecture/target-architecture.md", "Target cloud, data, resilience and security design"),
    ("docs/architecture/capacity-plan.md", "Capacity and scaling considerations"),
    ("docs/operations/deployment.md", "Current deployment readiness and operational caveats"),
], [2.8, 4.15])

add_heading(doc, "10. Technology Stack at a Glance", 1)
add_table(doc, ["Category", "Technologies"], [
    ("Frontend", "React 18, Vite 5, Axios, React Router, Socket.IO Client, Leaflet"),
    ("Backend", "Node.js, Express 4, Mongoose 7, Socket.IO, JWT, KafkaJS"),
    ("Data", "MongoDB, Redis, Kafka, ClickHouse, S3-compatible object storage"),
    ("Agents", "Python, psutil, watchdog, cryptography, YARA, pyudev; Java 17 / Android"),
    ("Detection", "Suricata, Zeek, YARA, VirusTotal, Velociraptor"),
    ("Security", "AES-256-GCM, TLS/mTLS, JWT, Argon2/bcrypt, signed payloads, integrity manifests"),
    ("Cloud target", "Load balancer/WAF, containers/Kubernetes, managed broker/data stores, KMS, secrets, monitoring"),
], [1.35, 5.6])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— End of document —")
r.italic = True
r.font.color.rgb = RGBColor.from_string(MID_GREY)

doc.core_properties.title = "AJNAT SOC Technology Stack"
doc.core_properties.subject = "Technology stack, cloud architecture and security overview"
doc.core_properties.author = "AJNAT SOC"
doc.core_properties.keywords = "AJNAT, SOC, technology stack, cloud, cybersecurity"
doc.save(OUTPUT)
print(OUTPUT)
